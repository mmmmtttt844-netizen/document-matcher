import streamlit as st
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from docx import Document
from docx.shared import Cm
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
import io
import math


# =========================================================
# إعداد الموقع
# =========================================================

st.set_page_config(
    page_title="PhotoPrint Pro",
    page_icon="📸",
    layout="centered"
)

st.title("📸 PhotoPrint Pro")
st.caption("قالب صور 3×4 — ست صور في السطر")


# =========================================================
# الإعدادات الثابتة
# =========================================================

PHOTO_WIDTH_CM = 3
PHOTO_HEIGHT_CM = 4

COLUMNS = 6

A4_WIDTH_CM = 21
A4_HEIGHT_CM = 29.7

TOP_MARGIN_CM = 0.5
SIDE_MARGIN_CM = 0.5

GAP_CM = 0.15


# =========================================================
# رفع الصورة
# =========================================================

uploaded = st.file_uploader(
    "📸 ارفع الصورة",
    type=["jpg", "jpeg", "png", "webp"]
)


# =========================================================
# قص الصورة إلى 3×4
# =========================================================

def crop_to_3x4(image):

    image = image.convert("RGB")

    target_ratio = 3 / 4
    current_ratio = image.width / image.height

    if current_ratio > target_ratio:

        new_width = int(
            image.height * target_ratio
        )

        left = (
            image.width - new_width
        ) // 2

        image = image.crop(
            (
                left,
                0,
                left + new_width,
                image.height
            )
        )

    else:

        new_height = int(
            image.width / target_ratio
        )

        top = (
            image.height - new_height
        ) // 2

        image = image.crop(
            (
                0,
                top,
                image.width,
                top + new_height
            )
        )

    return image


# =========================================================
# تحويل الصورة إلى JPG بالذاكرة
# =========================================================

def image_to_buffer(image):

    buffer = io.BytesIO()

    image.save(
        buffer,
        format="JPEG",
        quality=95
    )

    buffer.seek(0)

    return buffer


# =========================================================
# إنشاء PDF
# =========================================================

def create_pdf(image, copies):

    # تحويل إلى 3×4
    image = crop_to_3x4(image)

    image_buffer = image_to_buffer(image)

    image_reader = ImageReader(
        image_buffer
    )

    # A4 بالنقاط
    page_width, page_height = A4

    CM_TO_PT = 28.3464567

    photo_width = (
        PHOTO_WIDTH_CM * CM_TO_PT
    )

    photo_height = (
        PHOTO_HEIGHT_CM * CM_TO_PT
    )

    top_margin = (
        TOP_MARGIN_CM * CM_TO_PT
    )

    side_margin = (
        SIDE_MARGIN_CM * CM_TO_PT
    )

    gap = (
        GAP_CM * CM_TO_PT
    )

    # حساب عدد الأسطر
    rows = math.ceil(
        copies / COLUMNS
    )

    output = io.BytesIO()

    pdf = canvas.Canvas(
        output,
        pagesize=A4
    )

    current_copy = 0

    # عدد الصور التي يمكن وضعها بالطول
    available_height = (
        page_height
        - top_margin
        - side_margin
    )

    rows_per_page = int(
        (
            available_height + gap
        )
        /
        (
            photo_height + gap
        )
    )

    rows_per_page = max(
        1,
        rows_per_page
    )

    photos_per_page = (
        COLUMNS * rows_per_page
    )

    while current_copy < copies:

        page_start = current_copy

        # عدد الصور بهذه الصفحة
        page_remaining = (
            copies - current_copy
        )

        page_photos = min(
            page_remaining,
            photos_per_page
        )

        page_rows = math.ceil(
            page_photos / COLUMNS
        )

        for i in range(page_photos):

            row = i // COLUMNS
            column = i % COLUMNS

            x = (
                side_margin
                +
                column *
                (
                    photo_width + gap
                )
            )

            y = (
                page_height
                -
                top_margin
                -
                photo_height
                -
                row *
                (
                    photo_height + gap
                )
            )

            pdf.drawImage(
                image_reader,
                x,
                y,
                width=photo_width,
                height=photo_height,
                preserveAspectRatio=False
            )

            # -----------------------------
            # علامات القص
            # -----------------------------

            pdf.setLineWidth(0.3)

            # يسار
            pdf.line(
                x - 3,
                y,
                x - 1,
                y
            )

            pdf.line(
                x,
                y - 3,
                x,
                y - 1
            )

            # يمين
            pdf.line(
                x + photo_width + 1,
                y,
                x + photo_width + 3,
                y
            )

            pdf.line(
                x + photo_width,
                y - 3,
                x + photo_width,
                y - 1
            )

            # أعلى
            pdf.line(
                x - 3,
                y + photo_height,
                x - 1,
                y + photo_height
            )

            pdf.line(
                x,
                y + photo_height + 1,
                x,
                y + photo_height + 3
            )

            pdf.line(
                x + photo_width + 1,
                y + photo_height,
                x + photo_width + 3,
                y + photo_height
            )

            pdf.line(
                x + photo_width,
                y + photo_height + 1,
                x + photo_width,
                y + photo_height + 3
            )

        current_copy += page_photos

        pdf.showPage()

    pdf.save()

    output.seek(0)

    return output


# =========================================================
# إنشاء Word
# =========================================================

def create_word(image, copies):

    image = crop_to_3x4(image)

    image_buffer = image_to_buffer(image)

    # -----------------------------------------
    # إنشاء مستند Word
    # -----------------------------------------

    document = Document()

    section = document.sections[0]

    # حجم A4
    section.page_width = Cm(
        A4_WIDTH_CM
    )

    section.page_height = Cm(
        A4_HEIGHT_CM
    )

    # الهوامش
    section.top_margin = Cm(
        TOP_MARGIN_CM
    )

    section.bottom_margin = Cm(
        0.5
    )

    section.left_margin = Cm(
        SIDE_MARGIN_CM
    )

    section.right_margin = Cm(
        SIDE_MARGIN_CM
    )

    # -----------------------------------------
    # عدد الصفوف
    # -----------------------------------------

    rows = math.ceil(
        copies / COLUMNS
    )

    # إنشاء جدول 6 أعمدة
    table = document.add_table(
        rows=rows,
        cols=COLUMNS
    )

    table.autofit = False

    # -----------------------------------------
    # إضافة الصور
    # -----------------------------------------

    current = 0

    for row in table.rows:

        for cell in row.cells:

            cell.width = Cm(
                PHOTO_WIDTH_CM
            )

            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            if current < copies:

                paragraph = (
                    cell.paragraphs[0]
                )

                paragraph.alignment = 1

                run = paragraph.add_run()

                # إضافة الصورة
                run.add_picture(
                    io.BytesIO(
                        image_buffer.getvalue()
                    ),
                    width=Cm(
                        PHOTO_WIDTH_CM
                    ),
                    height=Cm(
                        PHOTO_HEIGHT_CM
                    )
                )

                current += 1

            else:

                cell.text = ""

    # -----------------------------------------
    # حفظ Word
    # -----------------------------------------

    output = io.BytesIO()

    document.save(
        output
    )

    output.seek(0)

    return output


# =========================================================
# تشغيل البرنامج
# =========================================================

if uploaded:

    image = Image.open(
        uploaded
    )

    st.success(
        "تم رفع الصورة بنجاح ✅"
    )

    # -----------------------------------------
    # معاينة
    # -----------------------------------------

    st.subheader(
        "👀 الصورة"
    )

    st.image(
        image,
        width=220
    )

    st.divider()

    # -----------------------------------------
    # عدد النسخ
    # -----------------------------------------

    copies = st.number_input(
        "🔢 عدد النسخ",
        min_value=1,
        max_value=200,
        value=6,
        step=1
    )

    st.info(
        f"📄 القالب: A4 | "
        f"📸 المقاس: 3×4 سم | "
        f"🖼️ 6 صور بالسطر | "
        f"🔢 عدد النسخ: {copies}"
    )

    # -----------------------------------------
    # زر إنشاء الملفات
    # -----------------------------------------

    if st.button(
        "🖨️ تجهيز الملفات",
        type="primary",
        use_container_width=True
    ):

        with st.spinner(
            "جاري تجهيز Word و PDF..."
        ):

            pdf_file = create_pdf(
                image,
                copies
            )

            word_file = create_word(
                image,
                copies
            )

        st.success(
            "✅ تم تجهيز الملفات بنجاح"
        )

        # -------------------------------------
        # أزرار التحميل
        # -------------------------------------

        col1, col2 = st.columns(2)

        with col1:

            st.download_button(
                "📄 تحميل Word",
                data=word_file,
                file_name="صور_3x4.docx",
                mime=(
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml.document"
                ),
                use_container_width=True
            )

        with col2:

            st.download_button(
                "📕 تحميل PDF",
                data=pdf_file,
                file_name="صور_3x4.pdf",
                mime="application/pdf",
                use_container_width=True
            )
